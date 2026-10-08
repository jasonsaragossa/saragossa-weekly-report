"""The Contract MBR as Jonny's own workbook."""
import io
import zipfile

import openpyxl

from shared import mbr_export as E
from shared import xlsx_cells as X


def _view(values):
    return {"values": values}


def _sheet(xlsx: bytes, name: str):
    return openpyxl.load_workbook(io.BytesIO(xlsx))[name]


def test_every_mapped_row_is_the_row_jonny_labelled_that_way():
    for template, sheet, rows in (("mbr_director.xlsx", "Contract Performance", E.DIRECTOR_ROWS),
                                  ("mbr_consultant.xlsx", "Performance", E.CONSULTANT_ROWS)):
        ws = openpyxl.load_workbook(E.TEMPLATES / template)[sheet]
        for row, (label, _key) in rows.items():
            assert (ws.cell(row, 1).value or "").strip() == label.strip(), (template, row)


def test_the_templates_carry_none_of_the_source_data():
    for template in ("mbr_director.xlsx", "mbr_consultant.xlsx"):
        z = zipfile.ZipFile(E.TEMPLATES / template)
        strings = z.read("xl/sharedStrings.xml").decode()
        for name in ("Louis Warren", "Peter Head", "Flynn Kennedy", "Lewis Fuller"):
            assert name not in strings, (template, name)
        assert "xl/calcChain.xml" not in z.namelist()


def test_a_months_figures_land_in_its_column_and_formulas_survive():
    team = _view({"P1": {"gp_month": 283251, "runners_split": 84.0}, "P9": {"wgp_running": 103689.83}})
    ws = _sheet(E.director(team, [("Louis Warren", _view({"P9": {"wgp_running": 25000}}))]),
                "Contract Performance")
    assert ws["C2"].value == 283251 and ws["C5"].value == 84
    assert ws["N10"].value == 103689.83
    assert ws["F2"].value == "=C2+D2+E2"          # Jonny's own quarter formula
    assert ws["D2"].value is None                 # nothing invented for an empty month


def test_the_director_wgp_tracker_has_a_row_per_consultant():
    team = _view({})
    xlsx = E.director(team, [("Louis Warren", _view({"P1": {"wgp_running": 11336.53}})),
                             ("Peter Head", _view({"P9": {"wgp_running": 21000.0}}))])
    ws = _sheet(xlsx, "WNF Tracker")
    assert ws["A4"].value == "Louis Warren" and ws["C4"].value == 11336.53
    assert ws["A5"].value == "Peter Head" and ws["N5"].value == 21000


def test_a_consultant_export_fills_their_sheet_and_tracker():
    xlsx = E.consultant(_view({"P2": {"wgp_running": 13298.91, "runners": 17, "deals": 2.0}}))
    ws = _sheet(xlsx, "Performance")
    assert ws["D2"].value == 13298.91 and ws["D3"].value == 17 and ws["D11"].value == 2
    assert ws["D4"].value == "=D2/D3"
    assert _sheet(xlsx, "WNF Tracker")["C4"].value == 13298.91


def test_excel_is_told_to_recalculate_on_opening():
    z = zipfile.ZipFile(io.BytesIO(E.consultant(_view({}))))
    assert 'fullCalcOnLoad="1"' in z.read("xl/workbook.xml").decode()


def test_set_cells_keeps_the_style_and_never_overwrites_a_formula():
    xml = ('<worksheet><sheetData><row r="2"><c r="A2" t="s"><v>0</v></c><c r="C2" s="7"/>'
           '<c r="F2" s="9"><f>C2+D2</f><v>5</v></c></row></sheetData></worksheet>')
    out = X.set_cells(xml, {"C2": 12.5, "F2": 99, "D2": 3, "B7": "Name"})
    assert '<c r="C2" s="7"><v>12.5</v></c>' in out
    assert "<f>C2+D2</f>" in out and "<v>99</v>" not in out
    assert out.index('r="C2"') < out.index('r="D2"') < out.index('r="F2"')     # column order kept
    assert '<row r="7"><c r="B7" t="inlineStr"><is><t>Name</t></is></c></row>' in out
